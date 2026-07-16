#!/usr/bin/env python3
"""Apply the VBX brand layer to a .docx.

Running header on every page: logo + wordmark (top-left) + descriptor and the
BUILD WHAT DOESN'T EXIST YET tagline (right), over a teal rule. Footer gets a
gold accent rule; existing footer text and page-number fields are preserved.
Body content is left untouched. Optionally strip disallowed strings.

Usage:
  python brand_docx.py IN.docx OUT.docx \
      --descriptor "Task Order Pricing Playbook" \
      [--logo /path/to/visionblox-logo.png] \
      [--strip "GSA MAS SIN 54151HEAL"]   # repeatable; removes "  |  <text>" too

Requires: python-docx
"""
import argparse, os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x23, 0x2D, 0x5A)
TEAL = RGBColor(0x2E, 0xA8, 0x91)
MUTED = RGBColor(0x88, 0x92, 0xA4)
TEAL_HEX, GOLD_HEX = "2EA891", "F7B801"
DEFAULT_LOGO = os.path.join(os.path.dirname(__file__), "..", "assets", "visionblox-logo.png")
TAGLINE = "B U I L D   W H A T   D O E S N ' T   E X I S T   Y E T"


def _cell_margins(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k in ("top", "start", "bottom", "end", "left", "right"):
        if k in kw:
            e = OxmlElement("w:" + k)
            e.set(qn("w:w"), str(kw[k])); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def _no_borders(table):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none"); e.set(qn("w:sz"), "0"); e.set(qn("w:space"), "0")
        b.append(e)
    table._tbl.tblPr.append(b)


def _border(p, edge, color, sz, space):
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    e = OxmlElement("w:" + edge)
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space)); e.set(qn("w:color"), color)
    pbdr.append(e)


def _clear(container):
    el = container._element
    for c in list(el):
        if c.tag in (qn("w:p"), qn("w:tbl")):
            el.remove(c)


def build_header(header, descriptor, logo):
    _clear(header)
    tbl = header.add_table(rows=1, cols=2, width=Inches(6.5))
    tbl.autofit = False
    _no_borders(tbl)
    left, right = tbl.rows[0].cells
    left.width, right.width = Inches(2.5), Inches(4.0)
    left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    for c in (left, right):
        _cell_margins(c, top=0, bottom=0, left=0, right=0)
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lp.paragraph_format.space_before = Pt(0); lp.paragraph_format.space_after = Pt(0)
    lp.add_run().add_picture(logo, width=Inches(1.7))
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = Pt(0); rp.paragraph_format.space_after = Pt(1)
    r = rp.add_run(descriptor)
    r.font.name = "Arial"; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = NAVY
    rp2 = right.add_paragraph()
    rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp2.paragraph_format.space_before = Pt(0); rp2.paragraph_format.space_after = Pt(0)
    t = rp2.add_run(TAGLINE)
    t.font.name = "Arial"; t.font.size = Pt(6.5); t.font.bold = True; t.font.color.rgb = TEAL
    div = header.add_paragraph()
    div.paragraph_format.space_before = Pt(2); div.paragraph_format.space_after = Pt(0)
    _border(div, "bottom", TEAL_HEX, 12, 1)


def style_footer(footer):
    if not footer.paragraphs:
        return
    p = footer.paragraphs[0]
    _border(p, "top", GOLD_HEX, 8, 4)
    for r in p.runs:
        r.font.name = "Arial"
        if r.font.size is None:
            r.font.size = Pt(7.5)
        r.font.color.rgb = MUTED


def strip_text(doc, phrases):
    def fix(paras):
        for p in paras:
            for r in p.runs:
                for ph in phrases:
                    if ph in r.text:
                        new = r.text.replace("  |  " + ph, "")
                        if ph in new:
                            new = new.replace(ph + "  |  ", "").replace(ph, "")
                        r.text = new
    fix(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                fix(c.paragraphs)
    for s in doc.sections:
        fix(s.header.paragraphs); fix(s.footer.paragraphs)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("infile"); ap.add_argument("outfile")
    ap.add_argument("--descriptor", required=True)
    ap.add_argument("--logo", default=DEFAULT_LOGO)
    ap.add_argument("--strip", action="append", default=[])
    a = ap.parse_args()
    d = Document(a.infile)
    for s in d.sections:
        s.different_first_page_header_footer = False
        build_header(s.header, a.descriptor, a.logo)
        style_footer(s.footer)
    if a.strip:
        strip_text(d, a.strip)
    d.save(a.outfile)
    print("branded ->", a.outfile)


if __name__ == "__main__":
    main()
