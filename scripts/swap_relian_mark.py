#!/usr/bin/env python3
"""
Replace the generic placeholder glyph on a Relian deck's title slide with the
real Relian "Ledger Mark".

Both Relian decks open the same way: a teal (#2EA891) disc with a generic stock
glyph dropped on top -- a layers icon on the Substrate Briefing, a shield on the
external capabilities deck -- standing in for a product mark that did not exist
yet. This removes the disc and the glyph together and places the real mark on the
navy background in their footprint, left-aligned to the slide's 0.60in margin and
scaled so the mark's ink is as tall as the disc it replaces.

The disc and glyph are located structurally (a square, teal-filled shape in the
upper left, plus the picture inside its bounds) rather than by index, so the
script fails loudly if a deck's layout is not what it expects.

Usage:
    python scripts/swap_relian_mark.py \
        --deck branded_docs/VBX_Relian_Substrate_Briefing.pptx \
        --mark branded_docs/assets/relian/relian_mark_dark.png
"""
import argparse

from pptx import Presentation
from pptx.util import Inches, Emu
from PIL import Image

DISC_FILL = "2EA891"        # vbx teal
MAX_DISC_IN = 1.40          # discs are ~1.0-1.15in; anything larger is not one
MAX_DISC_TOP_IN = 3.50      # the mark slot sits in the upper half of the slide
SLIDE_MARGIN_IN = 0.60


def _in(v):
    return Emu(v).inches if v is not None else 0.0


def _fill_hex(sh):
    try:
        if sh.fill.type == 1:
            return str(sh.fill.fore_color.rgb).upper()
    except Exception:
        pass
    return None


def find_disc(slide):
    """The teal, square placeholder disc on the title slide."""
    hits = []
    for sh in slide.shapes:
        if sh.left is None or sh.shape_type == 13:
            continue
        w, h = _in(sh.width), _in(sh.height)
        if (_fill_hex(sh) == DISC_FILL and abs(w - h) < 0.02
                and w <= MAX_DISC_IN and _in(sh.top) <= MAX_DISC_TOP_IN):
            hits.append(sh)
    if len(hits) != 1:
        raise SystemExit(f"expected exactly 1 placeholder disc, found {len(hits)}")
    return hits[0]


def find_glyph(slide, disc):
    """The picture sitting inside the disc's bounds."""
    dl, dt, dr, db = (_in(disc.left), _in(disc.top),
                      _in(disc.left) + _in(disc.width), _in(disc.top) + _in(disc.height))
    hits = [sh for sh in slide.shapes
            if sh.shape_type == 13 and sh.left is not None
            and dl <= _in(sh.left) and dt <= _in(sh.top)
            and _in(sh.left) + _in(sh.width) <= dr + 0.01
            and _in(sh.top) + _in(sh.height) <= db + 0.01]
    if len(hits) != 1:
        raise SystemExit(f"expected exactly 1 glyph inside the disc, found {len(hits)}")
    return hits[0]


def swap(deck, mark_png, out):
    prs = Presentation(deck)
    slide = prs.slides[0]

    disc = find_disc(slide)
    glyph = find_glyph(slide, disc)
    height = _in(disc.height)
    top = _in(disc.top)

    iw, ih = Image.open(mark_png).size          # pre-cropped to the ink bounds
    width = height * iw / ih

    # Take the disc's exact place in the z-order. Appending instead would be
    # harmless here, but inserting at a fixed low index would bury the mark
    # behind the full-slide background rectangle these decks draw first.
    tree = slide.shapes._spTree
    z = list(tree).index(disc._element)
    glyph_rid = glyph._element.blip_rId
    for sh in (glyph, disc):
        sh._element.getparent().remove(sh._element)
    # Removing the shape leaves the relationship -- and so the generic glyph
    # itself -- inside the package. Drop it, or the icon we just "replaced" is
    # still shipped in the .pptx.
    slide.part.drop_rel(glyph_rid)

    pic = slide.shapes.add_picture(mark_png, Inches(SLIDE_MARGIN_IN), Inches(top),
                                   Inches(width), Inches(height))
    tree.remove(pic._element)
    tree.insert(z, pic._element)

    prs.save(out)
    print(f"{deck}: disc {height:.2f}in + glyph removed -> mark "
          f"{width:.2f}x{height:.2f}in at ({SLIDE_MARGIN_IN}, {top:.2f})")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", required=True)
    ap.add_argument("--mark", required=True)
    ap.add_argument("--out")
    a = ap.parse_args()
    swap(a.deck, a.mark, a.out or a.deck)


if __name__ == "__main__":
    main()
