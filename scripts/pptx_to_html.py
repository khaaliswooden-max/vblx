#!/usr/bin/env python3
"""
Approximate .pptx -> HTML slide renderer, so a headless-Chromium screenshot can
serve as a visual proof (LibreOffice cannot load files in this environment).

Covers the shape vocabulary this deck uses: rect / roundRect / ellipse, solid
fills, hairline outlines, outer shadows, pictures, and text frames with runs,
alignment, vertical anchor, bullets, letter-spacing and paragraph spacing.
Arial maps to Liberation Sans, which is metric-compatible, so text wrapping and
overflow in the render match what PowerPoint will do.
"""
import base64
import html
import sys

from pptx import Presentation
from pptx.util import Emu

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
EMU_IN = 914400.0
PX = 96.0  # px per inch


def px(emu):
    return (emu or 0) / EMU_IN * PX


def find(el, path):
    return el.find(path) if el is not None else None


def css_color(c):
    """Turn an <a:srgbClr> element into a CSS color, honouring <a:alpha>."""
    if c is None:
        return None
    val = c.get("val")
    if not val:
        return None
    alpha = c.find(A + "alpha")
    if alpha is None:
        return "#" + val
    a = int(alpha.get("val")) / 100000.0
    r, g, b = (int(val[i:i + 2], 16) for i in (0, 2, 4))
    return f"rgba({r},{g},{b},{a:.3f})"


def solid_hex(parent):
    """Color of a direct a:solidFill child (fills, lines), or None."""
    if parent is None:
        return None
    return css_color(find(parent.find(A + "solidFill"), A + "srgbClr"))


def shape_css(sp):
    """Fill / border / radius / shadow CSS for an autoshape."""
    spPr = sp._element.find(".//" + A + "prstGeom")
    spPr_el = sp._element.spPr if hasattr(sp._element, "spPr") else None
    css = []
    geom = sp._element.find(".//" + A + "prstGeom")
    prst = geom.get("prst") if geom is not None else "rect"

    node = None
    for child in sp._element:
        if child.tag.endswith("}spPr"):
            node = child
            break
    fill = solid_hex(node)
    if fill:
        css.append(f"background:{fill}")
    ln = node.find(A + "ln") if node is not None else None
    if ln is not None:
        lc = solid_hex(ln)
        w = int(ln.get("w") or 9525)
        if lc:
            css.append(f"border:{max(1, round(px(w)))}px solid {lc}")
            css.append("box-sizing:border-box")
    if prst == "roundRect":
        adj = 0.16667
        gd = geom.find(A + "avLst/" + A + "gd") if geom is not None else None
        if gd is not None and gd.get("fmla", "").startswith("val "):
            adj = int(gd.get("fmla").split()[1]) / 100000.0
        w, h = px(sp.width), px(sp.height)
        css.append(f"border-radius:{min(w, h) * adj:.1f}px")
    elif prst == "ellipse":
        css.append("border-radius:50%")
    eff = node.find(A + "effectLst") if node is not None else None
    if eff is not None:
        sh = eff.find(A + "outerShdw")
        if sh is not None:
            blur = px(int(sh.get("blurRad") or 0))
            dist = px(int(sh.get("dist") or 0))
            # DrawingML nests the shadow color directly under a:outerShdw --
            # there is no a:solidFill wrapper here, unlike fills and lines.
            col = css_color(sh.find(A + "srgbClr")) or "rgba(0,0,0,0.13)"
            css.append(f"box-shadow:0 {dist:.1f}px {blur:.1f}px {col}")
    return ";".join(css)


def para_html(p, default_align):
    pPr = p._p.find(A + "pPr")
    align = default_align
    marL = indent = 0
    bullet = None
    space_after = 0
    if pPr is not None:
        align = {"ctr": "center", "r": "right", "l": "left"}.get(pPr.get("algn"), align)
        marL = px(int(pPr.get("marL") or 0))
        indent = px(int(pPr.get("indent") or 0))
        bu = pPr.find(A + "buChar")
        if bu is not None:
            bullet = bu.get("char")
        sa = pPr.find(A + "spcAft/" + A + "spcPts")
        if sa is not None:
            space_after = int(sa.get("val")) / 100.0 * PX / 72.0
    runs = []
    for r in p.runs:
        rPr = r._r.find(A + "rPr")
        st = []
        sz = 18.0
        if rPr is not None:
            if rPr.get("sz"):
                sz = int(rPr.get("sz")) / 100.0
            if rPr.get("b") == "1":
                st.append("font-weight:700")
            if rPr.get("i") == "1":
                st.append("font-style:italic")
            if rPr.get("spc"):
                st.append(f"letter-spacing:{int(rPr.get('spc')) / 100.0 * PX / 72.0:.2f}px")
            col = solid_hex(rPr)
            if col:
                st.append(f"color:{col}")
            latin = rPr.find(A + "latin")
            fam = latin.get("typeface") if latin is not None else "Arial"
            st.append(f"font-family:'{fam}','Liberation Sans',sans-serif")
        st.append(f"font-size:{sz * PX / 72.0:.2f}px")
        st.append(f"line-height:{sz * PX / 72.0 * 1.21:.2f}px")
        runs.append(f"<span style=\"{';'.join(st)}\">{html.escape(r.text)}</span>")
    body = "".join(runs) or "&nbsp;"
    style = [f"text-align:{align}", f"margin-bottom:{space_after:.1f}px"]
    if bullet:
        # bullet sits at marL+indent, text at marL -- a two-column flex row
        style.append(f"padding-left:{marL + indent:.1f}px")
        style.append("display:flex")
        body = (f"<span style=\"flex:none;width:{-indent:.1f}px\">"
                f"{html.escape(bullet)}</span><span style='flex:1'>{body}</span>")
    elif marL:
        style.append(f"padding-left:{marL:.1f}px")
    return f"<p style=\"{';'.join(style)}\">{body}</p>"


def text_html(sp):
    tf = sp.text_frame
    bodyPr = sp._element.find(".//" + A + "bodyPr")
    anchor = (bodyPr.get("anchor") if bodyPr is not None else "t") or "t"
    justify = {"t": "flex-start", "ctr": "center", "b": "flex-end"}.get(anchor, "flex-start")
    def ins(name, dflt):
        v = bodyPr.get(name) if bodyPr is not None else None
        return px(int(v) if v is not None else dflt)
    pad = (f"{ins('tIns', 45720):.1f}px {ins('rIns', 91440):.1f}px "
           f"{ins('bIns', 45720):.1f}px {ins('lIns', 91440):.1f}px")
    default_align = "left"
    paras = "".join(para_html(p, default_align) for p in tf.paragraphs)
    return (f"<div class=tf style=\"justify-content:{justify};padding:{pad}\">"
            f"<div style='width:100%'>{paras}</div></div>")


def slide_html(slide, idx):
    bg = "#FFFFFF"
    for c in slide.element.iter(A + "srgbClr"):
        n = c.getparent()
        while n is not None:
            if n.tag.endswith("}bgPr"):
                bg = "#" + c.get("val")
                n = None
                break
            n = n.getparent()
        if bg != "#FFFFFF":
            break
    out = [f"<div class=slide id=s{idx} style=\"background:{bg}\">"]
    for sp in slide.shapes:
        if sp.left is None:
            continue
        pos = (f"left:{px(sp.left):.1f}px;top:{px(sp.top):.1f}px;"
               f"width:{px(sp.width):.1f}px;height:{px(sp.height):.1f}px")
        if sp.shape_type == 13:  # picture
            blob = base64.b64encode(sp.image.blob).decode()
            out.append(f"<img style=\"position:absolute;{pos}\" "
                       f"src=\"data:{sp.image.content_type};base64,{blob}\">")
            continue
        css = shape_css(sp)
        out.append(f"<div data-name=\"{html.escape(sp.name)}\" "
                   f"style=\"position:absolute;{pos};{css}\">")
        if sp.has_text_frame and sp.text_frame.text.strip():
            out.append(text_html(sp))
        out.append("</div>")
    out.append(f"<div class=badge>{idx}</div></div>")
    return "".join(out)


def main(src, prefix):
    """Write one HTML file per slide: <prefix>_1.html ... so each screenshot is exact."""
    prs = Presentation(src)
    W, H = px(prs.slide_width), px(prs.slide_height)
    head = f"""<meta charset=utf-8><style>
*{{margin:0;padding:0}} body{{background:#fff}}
.slide{{position:relative;width:{W:.0f}px;height:{H:.0f}px;overflow:hidden;
  font-family:'Liberation Sans',Arial,sans-serif}}
.tf{{position:absolute;inset:0;display:flex;flex-direction:column;overflow:visible}}
.badge{{display:none}}
</style>"""
    # Reports text that no longer fits its shape (Arial is wider than Calibri).
    probe = """<pre id=report style="position:absolute;left:0;top:2000px"></pre>
<script>
var out=[];
document.querySelectorAll('[data-name]').forEach(function(d){
  var tf=d.querySelector(':scope > .tf'); if(!tf) return;
  var inner=tf.firstElementChild;
  var need=inner.getBoundingClientRect().height
         + parseFloat(getComputedStyle(tf).paddingTop)
         + parseFloat(getComputedStyle(tf).paddingBottom);
  var have=d.getBoundingClientRect().height;
  if(need-have>1.5) out.push(d.dataset.name+' needs '+need.toFixed(1)
      +'px has '+have.toFixed(1)+'px :: '+tf.innerText.slice(0,60).replace(/\\n/g,' '));
});
document.getElementById('report').textContent=out.join('\\n')||'OK';
</script>"""
    for i, s in enumerate(prs.slides, 1):
        open(f"{prefix}_{i}.html", "w").write(head + slide_html(s, i) + probe)
    print("wrote", len(prs.slides), "slides", f"{W:.0f}x{H:.0f}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
