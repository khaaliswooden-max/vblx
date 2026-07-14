#!/usr/bin/env python3
"""
Apply the VisionBlox / VBX brand system to the SQ-OS pitch decks (.pptx).

What it does to every slide layout:
  * Embeds the real visionblox logo + wordmark (no more typed "VISIONBLOX" text):
      - dark (navy) title/closing slides  -> knockout logo (off-white wordmark) top-left
      - light content slides / dark content slides -> logo in the footer, bottom-left
  * Adds a gold (#F7B801) footer accent rule on content slides, matching the
    VBX docx branding convention (teal accents top, gold rule bottom).
  * Adds the "BUILD WHAT DOESN'T EXIST YET" tagline to hero slides.
  * Preserves all authored content, the VBX palette already in use, and the
    INTERNAL / CONFIDENTIAL banners.

The knockout logo (light wordmark on dark) is generated from the source logo:
the geometric mark and its teal accent are kept; the dark slate wordmark is
recolored off-white so it reads on the navy background.

Usage:
    python scripts/brand_sqos_decks.py \
        --logo visionblox-logo.png \
        --external SQ-OS_External.pptx \
        --internal SQ-OS_Internal.pptx \
        --outdir branded_docs
"""
import argparse, os, re, colorsys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ---- VBX palette (from tailwind.config.ts) --------------------------------
TEAL = RGBColor(0x2E, 0xA8, 0x91)
GOLD = RGBColor(0xF7, 0xB8, 0x01)
OFFWHITE = (245, 245, 240)
DARK_BGS = ("232D5A", "1A2140", "1B2347")
WORDMARK_SPLIT_Y = 255           # rows below this are the "visionblox" wordmark


def make_knockout(src_logo, dst):
    """Recolor the dark wordmark to off-white; keep the colorful mark + teal accent."""
    src = Image.open(src_logo).convert("RGBA")
    W, H = src.size
    px = src.load()
    ko = src.copy()
    k = ko.load()
    for y in range(WORDMARK_SPLIT_Y, H):
        for x in range(W):
            r, g, b, a = px[x, y]
            if a < 8:
                continue
            h, s, v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
            is_teal_green = (120 <= h * 360 <= 185) and s > 0.30
            if not is_teal_green:
                k[x, y] = (OFFWHITE[0], OFFWHITE[1], OFFWHITE[2], a)
    ko.save(dst)
    ratio = W / float(H)
    return dst, ratio


def _text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def _remove(sh):
    sh._element.getparent().remove(sh._element)


def _bg_hex(slide):
    m = re.search(r"<p:bg>.*?</p:bg>", slide.element.xml, re.S)
    if m:
        c = re.search(r'srgbClr val="([0-9A-Fa-f]{6})"', m.group(0))
        if c:
            return c.group(1).upper()
    return "FFFFFF"


def _add_logo(slide, path, left, top, width, ratio):
    return slide.shapes.add_picture(path, Inches(left), Inches(top),
                                    Inches(width), Inches(width / ratio))


def _add_rule(slide, left, top, width, height, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _add_tagline(slide, left, top, width):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = "BUILD WHAT DOESN'T EXIST YET"
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = TEAL
    return tb


def brand(src_pptx, color_logo, ko_logo, ratio, hero_idxs, out_pptx):
    prs = Presentation(src_pptx)
    for i, s in enumerate(prs.slides):
        dark = _bg_hex(s) in DARK_BGS
        logo = ko_logo if dark else color_logo
        if i in hero_idxs:
            for sh in list(s.shapes):
                if _text(sh) == "VISIONBLOX":
                    _remove(sh)
            for sh in list(s.shapes):                    # nudge eyebrow below lockup
                t = _text(sh)
                if (t.startswith("ZUUP INNOVATION LAB") or t == "LET'S TALK") \
                        and Emu(sh.top).inches < 1.5:
                    sh.top = Inches(1.55)
            _add_logo(s, logo, 0.7, 0.5, 2.25, ratio)
            _add_tagline(s, 8.2, 0.72, 4.43)
        else:
            for sh in list(s.shapes):                    # drop old typed footer wordmark
                t = _text(sh)
                if t.startswith("VISIONBLOX") and "ZUUP" in t and "·" in t:
                    _remove(sh)
            for sh in list(s.shapes):                    # lift captions out of footer band
                if not sh.has_text_frame:
                    continue
                top = Emu(sh.top).inches
                h = Emu(sh.height).inches
                w = Emu(sh.width).inches
                if w > 2.0 and (top + h) > 6.86:
                    sh.top = Inches(max(0.0, 6.82 - h))
            _add_rule(s, 0.7, 6.96, 11.93, 0.028, GOLD)
            _add_logo(s, logo, 0.7, 7.0, 1.05, ratio)
    prs.save(out_pptx)
    return out_pptx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--logo", default="visionblox-logo.png")
    ap.add_argument("--external", required=True)
    ap.add_argument("--internal", required=True)
    ap.add_argument("--outdir", default="branded_docs")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    ko_path = os.path.join(a.outdir, "visionblox-logo-knockout.png")
    _, ratio = make_knockout(a.logo, ko_path)
    ext = brand(a.external, a.logo, ko_path, ratio, {0, 7},
                os.path.join(a.outdir, "VBX_SQ-OS_External.pptx"))
    intr = brand(a.internal, a.logo, ko_path, ratio, {0},
                 os.path.join(a.outdir, "VBX_SQ-OS_Internal.pptx"))
    print("wrote:", ext)
    print("wrote:", intr)
    print("wrote:", ko_path)


if __name__ == "__main__":
    main()
