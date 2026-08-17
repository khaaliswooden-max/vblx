#!/usr/bin/env python3
"""
Apply the VisionBlox / VBX brand system to the Relian(TM) external capabilities
deck (.pptx).

Companion to scripts/brand_sqos_decks.py; adapted for a deck that draws its own
placeholder mark (four squares in a rounded container) next to a *typed*
"VISIONBLOX" wordmark. The VBX brand system forbids typing the wordmark, so:

  * Dark hero slides (title / closing) -> the drawn mark + typed wordmark are
    removed and replaced with the real KNOCKOUT lockup top-left, and the
    "BUILD WHAT DOESN'T EXIST YET" tagline is placed top-right in teal.
  * Light content slides -> the drawn footer mark + typed wordmark are removed
    and replaced with the real FULL-COLOR lockup bottom-left, sitting on the
    gold (#F7B801) footer accent rule. Page numbers, the "Relian(TM) .
    Visionblox LLC" footer and all authored content are preserved.
  * Dark content slides and the references slide carry no footer in the source;
    they get the same gold rule + lockup (knockout on dark) so every slide
    carries the mark.
  * Off-token colors are normalized to VBX palette values (see COLOR_MAP).

Usage:
    python scripts/brand_relian_deck.py \
        --src <deck>.pptx \
        --out branded_docs/VBX_Relian_Capabilities_External.pptx
"""
import argparse
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ---- VBX palette (tailwind.config.ts / references/palette.md) -------------
TEAL = RGBColor(0x2E, 0xA8, 0x91)
GOLD = RGBColor(0xF7, 0xB8, 0x01)

# Off-token -> VBX token. Applied to shape fills, table cell fills and runs.
COLOR_MAP = {
    "1B2247": "1B2347",   # near-navy-dark   -> vbx navy-dark
    "5A6373": "8892A4",   # generic gray ink -> vbx muted
}

SKILL_ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            os.pardir, ".claude", "skills", "vbx-branding", "assets")
LOGO_COLOR = os.path.normpath(os.path.join(SKILL_ASSETS, "visionblox-logo.png"))
LOGO_KO = os.path.normpath(os.path.join(SKILL_ASSETS, "visionblox-logo-knockout.png"))

WORDMARK = "VISIONBLOX"
HERO_LOGO_W = 2.2          # in
FOOTER_LOGO_W = 1.0        # in
RULE_TOP = 6.96            # in
RULE_H = 0.028             # in
FOOTER_LOGO_TOP = 6.99     # in


def _in(v):
    return Emu(v).inches if v is not None else 0.0


def _text(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def _remove(sh):
    sh._element.getparent().remove(sh._element)


def _fill_hex(sh):
    try:
        if sh.fill.type == 1:
            return str(sh.fill.fore_color.rgb).upper()
    except Exception:
        pass
    return None


def _is_dark(slide):
    """This deck paints its background with a full-bleed rectangle, not <p:bg>."""
    for sh in slide.shapes:
        if sh.width is None or sh.height is None:
            continue
        if _in(sh.width) > 13.0 and _in(sh.height) > 7.0:
            hexv = _fill_hex(sh)
            if hexv and hexv in ("232D5A", "1B2247", "1B2347", "2A3560", "1A2140"):
                return True
    return False


def normalize_colors(slide):
    """Remap off-token hexes to VBX tokens (fills, table cells, run colors)."""
    n = 0
    for sh in slide.shapes:
        cur = _fill_hex(sh)
        if cur in COLOR_MAP:
            sh.fill.fore_color.rgb = RGBColor.from_string(COLOR_MAP[cur])
            n += 1
        if sh.has_text_frame:
            n += _normalize_tf(sh.text_frame)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    try:
                        if cell.fill.type == 1:
                            c = str(cell.fill.fore_color.rgb).upper()
                            if c in COLOR_MAP:
                                cell.fill.fore_color.rgb = RGBColor.from_string(COLOR_MAP[c])
                                n += 1
                    except Exception:
                        pass
                    n += _normalize_tf(cell.text_frame)
    return n


def _normalize_tf(tf):
    n = 0
    for para in tf.paragraphs:
        for run in para.runs:
            try:
                cur = str(run.font.color.rgb).upper()
            except Exception:
                continue
            if cur in COLOR_MAP:
                run.font.color.rgb = RGBColor.from_string(COLOR_MAP[cur])
                n += 1
    return n


def strip_drawn_lockup(slide):
    """Remove the typed VISIONBLOX wordmark and the small squares drawn beside it.

    Returns (left, top) of the removed lockup, or None if the slide had none.
    """
    marks = [sh for sh in slide.shapes if _text(sh) == WORDMARK]
    if not marks:
        return None
    anchor = None
    for word in marks:
        w_left, w_top, w_bot = _in(word.left), _in(word.top), _in(word.top) + _in(word.height)
        band_top, band_bot = w_top - 0.25, w_bot + 0.25
        for sh in list(slide.shapes):
            if sh is word or sh.has_text_frame and _text(sh):
                continue
            if sh.width is None or sh.height is None:
                continue
            # small decorative glyph shapes sitting to the LEFT of the wordmark
            if _in(sh.width) <= 0.55 and _in(sh.height) <= 0.55 \
                    and _in(sh.left) < w_left \
                    and band_top <= _in(sh.top) <= band_bot:
                anchor = (_in(sh.left), _in(sh.top)) if anchor is None else \
                    (min(anchor[0], _in(sh.left)), min(anchor[1], _in(sh.top)))
                _remove(sh)
        anchor = (w_left, w_top) if anchor is None else anchor
        _remove(word)
    return anchor


def add_logo(slide, path, left, top, width):
    img = Image.open(path)
    ratio = img.size[0] / float(img.size[1])
    return slide.shapes.add_picture(path, Inches(left), Inches(top),
                                    Inches(width), Inches(width / ratio))


def add_rule(slide, left, top, width, height, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def set_tagline(slide, existing=None, left=8.3, top=0.72, width=4.43):
    """Place the tagline top-right in teal, reusing the authored one if present."""
    if existing is not None:
        existing.left, existing.top, existing.width = Inches(left), Inches(top), Inches(width)
        for para in existing.text_frame.paragraphs:
            para.alignment = PP_ALIGN.RIGHT
        return existing
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = "BUILD WHAT DOESN’T EXIST YET"
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = TEAL
    r._r.get_or_add_rPr().set("spc", "300")   # brand letter-spacing
    return tb


def find_tagline(slide):
    for sh in slide.shapes:
        if _text(sh).upper().startswith("BUILD WHAT DOESN"):
            return sh
    return None


def existing_footer_rule(slide):
    """The authored hairline rule across the footer, if any."""
    for sh in slide.shapes:
        if sh.width is None or sh.height is None:
            continue
        if _in(sh.height) <= 0.05 and _in(sh.width) > 10.0 and _in(sh.top) > 6.5:
            return sh
    return None


def brand(src, out, hero_idxs):
    prs = Presentation(src)
    slide_w = _in(prs.slide_width)
    report = []
    for i, slide in enumerate(prs.slides):
        n = i + 1
        dark = _is_dark(slide)
        logo = LOGO_KO if dark else LOGO_COLOR
        recolored = normalize_colors(slide)
        stripped = strip_drawn_lockup(slide)
        notes = []
        if i in hero_idxs:
            add_logo(slide, logo, 0.6, 0.5, HERO_LOGO_W)
            set_tagline(slide, find_tagline(slide),
                        left=slide_w - 0.6 - 4.43, top=0.72, width=4.43)
            notes.append("hero: knockout lockup + tagline")
        else:
            rule = existing_footer_rule(slide)
            if rule is not None:
                rule.top = Inches(RULE_TOP)
                rule.height = Inches(RULE_H)
                rule.fill.solid()
                rule.fill.fore_color.rgb = GOLD
                notes.append("gold rule (reused)")
            else:
                add_rule(slide, 0.6, RULE_TOP, slide_w - 1.2, RULE_H, GOLD)
                notes.append("gold rule (added)")
            add_logo(slide, logo, 0.6, FOOTER_LOGO_TOP, FOOTER_LOGO_W)
            notes.append("footer lockup (%s)" % ("knockout" if dark else "full-color"))
        report.append("slide %-2d %-6s | %-2d recolored | wordmark %s | %s"
                      % (n, "dark" if dark else "light", recolored,
                         "replaced" if stripped else "n/a", "; ".join(notes)))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    return out, report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", default="branded_docs/VBX_Relian_Capabilities_External.pptx")
    ap.add_argument("--hero", default="1,12", help="1-based slide numbers treated as hero/dark title slides")
    a = ap.parse_args()
    hero = {int(x) - 1 for x in a.hero.split(",") if x.strip()}
    out, report = brand(a.src, a.out, hero)
    print("\n".join(report))
    print("wrote:", out)


if __name__ == "__main__":
    main()
