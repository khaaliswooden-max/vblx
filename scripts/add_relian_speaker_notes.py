#!/usr/bin/env python3
"""
Write the Relian(TM) external capabilities deck's speaker notes into each
slide's Notes pane.

The deck ships with 13 notes slides that are empty scaffolding (a notes body
placeholder holding a single empty run, plus the slide-number field), so the
notes only existed as prose outside the file. This fills them in from the
markdown source of truth, which stays readable for an email, a run-of-show
doc, or a teleprompter.

Notes source format (branded_docs/source/Relian_Capabilities_External_SpeakerNotes.md):

    **Slide 1 — Cover (…)**
    <note body, one paragraph>

Usage:
    python scripts/add_relian_speaker_notes.py \
        --deck branded_docs/VBX_Relian_Capabilities_External.pptx \
        --notes branded_docs/source/Relian_Capabilities_External_SpeakerNotes.md
"""
import argparse
import re

from pptx import Presentation
from pptx.util import Pt

HEADER = re.compile(r"^\*\*Slide\s+(\d+)\s*[—-].*\*\*\s*$")


def parse_notes(path):
    """-> {slide_number: note_text}. Body is every line until the next header."""
    notes, cur, buf = {}, None, []
    for line in open(path, encoding="utf-8").read().splitlines():
        m = HEADER.match(line.strip())
        if m:
            if cur is not None:
                notes[cur] = "\n".join(buf).strip()
            cur, buf = int(m.group(1)), []
        elif cur is not None:
            buf.append(line)
    if cur is not None:
        notes[cur] = "\n".join(buf).strip()
    return {k: v for k, v in notes.items() if v}


def apply_notes(deck, notes, out, font="Arial", size=12):
    prs = Presentation(deck)
    n_slides = len(prs.slides)
    missing = [i for i in range(1, n_slides + 1) if i not in notes]
    extra = [i for i in notes if i > n_slides]
    if missing or extra:
        raise SystemExit("notes/slide mismatch — missing %s, extra %s" % (missing, extra))
    for i, slide in enumerate(prs.slides, 1):
        tf = slide.notes_slide.notes_text_frame
        tf.clear()
        paras = [p for p in notes[i].split("\n\n") if p.strip()]
        for j, text in enumerate(paras):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = text.strip()
            run.font.name = font
            run.font.size = Pt(size)
    prs.save(out)
    return out, n_slides


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default="branded_docs/VBX_Relian_Capabilities_External.pptx")
    ap.add_argument("--notes", default="branded_docs/source/Relian_Capabilities_External_SpeakerNotes.md")
    ap.add_argument("--out", help="defaults to updating --deck in place")
    a = ap.parse_args()
    notes = parse_notes(a.notes)
    out, n = apply_notes(a.deck, notes, a.out or a.deck)
    print("wrote notes for %d/%d slides -> %s" % (len(notes), n, out))


if __name__ == "__main__":
    main()
